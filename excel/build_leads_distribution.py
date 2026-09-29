import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

OUT = sys.argv[1]
MAXR = int(__import__("os").environ.get("MAXR", 5000))
LAST = MAXR + 1
REPS = ["בן", "אור", "רותם", "בר", "אוריה", "אודל", "סדי", "מרים"]
REP_SLOTS = 20

P = "'הדבקת לידים'"
S = "'הגדרות'"
H = "'עזר'"
RES = "'תוצאה'"
SEED = f"{P}!$F$2"

F = "Arial"
font = Font(name=F, size=11)
bold = Font(name=F, size=11, bold=True)
hdr_font = Font(name=F, size=12, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
input_fill = PatternFill("solid", fgColor="FFF2CC")
panel_fill = PatternFill("solid", fgColor="DDEBF7")
gray_fill = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="A6A6A6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)

wb = Workbook()
wp = wb.active
wp.title = "הדבקת לידים"
wr = wb.create_sheet("תוצאה")
ws = wb.create_sheet("הגדרות")
wh = wb.create_sheet("עזר")
for w in (wp, wr, ws, wh):
    w.sheet_view.rightToLeft = True


def hdr(c, text):
    c.value = text
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = box


# ---------------- הגדרות ----------------
hdr(ws["A1"], "שמות נציגים")
for i in range(REP_SLOTS):
    c = ws.cell(row=i + 2, column=1, value=REPS[i] if i < len(REPS) else None)
    c.fill = input_fill
    c.font = font
    c.border = box
    c.alignment = right
hdr(ws["D1"], "פרמטר")
hdr(ws["E1"], "ערך")
params = [
    (2, "השלמת אפסים מובילים למספר זיהוי מספרי – אורך בספרות (0 = ללא השלמה)", 9, True),
    (3, "השלמת אפסים מובילים לטלפון מספרי – אורך בספרות (0 = ללא השלמה)", 10, True),
    (5, "מספר לידים שזוהו (מחושב)", f"=COUNT({H}!$E$2:$E${LAST})", False),
    (6, "מספר נציגים (מחושב)", "=COUNTA($A$2:$A$21)", False),
    (7, "לידים בסיסיים לכל נציג (מחושב)", "=IF(E6=0,0,INT(E5/E6))", False),
    (8, "נציגים שמקבלים ליד נוסף (מחושב)", "=IF(E6=0,0,MOD(E5,E6))", False),
    (9, "היסט סבב נציגים (מחושב)", f"=IF(E6=0,0,MOD(ABS(INT({SEED}))+E5,E6))", False),
]
for r, label, val, is_input in params:
    a = ws.cell(row=r, column=4, value=label)
    a.font = font
    a.alignment = right
    a.border = box
    b = ws.cell(row=r, column=5, value=val)
    b.font = font
    b.border = box
    b.alignment = center
    b.fill = input_fill if is_input else gray_fill
ws["D11"] = ("הערות: שמות הנציגים נכתבים בעמודה A ברצף, ללא שורות ריקות באמצע (עד 20 נציגים). "
             "השלמת אפסים חלה רק על ערכים שהודבקו כמספר (ואיבדו את האפס המוביל) – למשל 12345678 יוצג כ-012345678, "
             "ו-501234567 יוצג כ-0501234567. ערכים שהודבקו כטקסט נשמרים בדיוק כפי שהם. "
             "תאים בצהוב = ניתנים לעריכה; תאים באפור = מחושבים.")
ws["D11"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
ws["D11"].font = Font(name=F, size=10, italic=True)
ws.merge_cells("D11:E16")
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 3
ws.column_dimensions["C"].width = 3
ws.column_dimensions["D"].width = 62
ws.column_dimensions["E"].width = 12
dv_len = DataValidation(type="whole", operator="between", formula1="0", formula2="20",
                        showErrorMessage=True, errorTitle="ערך לא תקין", error="יש להזין מספר שלם בין 0 ל-20")
ws.add_data_validation(dv_len)
dv_len.add("E2:E3")

N, R, Q, REM, OFF = (f"{S}!$E$5", f"{S}!$E$6", f"{S}!$E$7", f"{S}!$E$8", f"{S}!$E$9")

# ---------------- עזר ----------------
helper_headers = ["מספר זיהוי מנורמל", "טלפון מנורמל", "שורה תקינה", "גיבוב 1", "ציון ערבוב",
                  "מיקום בערבוב", "זיהוי כפול", "שורת תוצאה", "סבב נציג", "מקום בתוך הנציג",
                  "מיקום בערבוב לשורה", "שורת מקור", "שם נציג"]
for j, t in enumerate(helper_headers, 1):
    hdr(wh.cell(row=1, column=j), t)
for i in range(2, LAST + 1):
    a, b = f"{P}!A{i}", f"{P}!B{i}"
    wh[f"A{i}"] = (f'=IF({a}="","",IF(ISNUMBER({a}),IF({S}!$E$2>0,TEXT({a},REPT("0",{S}!$E$2)),TEXT({a},"0")),TRIM({a})))')
    wh[f"B{i}"] = (f'=IF({b}="","",IF(ISNUMBER({b}),IF({S}!$E$3>0,TEXT({b},REPT("0",{S}!$E$3)),TEXT({b},"0")),TRIM({b})))')
    wh[f"C{i}"] = f'=IF(OR(A{i}<>"",B{i}<>""),1,0)'
    wh[f"D{i}"] = (f'=IF(C{i}=1,MOD(ABS(SIN(ROW()*12.9898+{SEED}*78.233'
                   f'+IFERROR(--RIGHT(A{i},6),LEN(A{i}))*0.0137+IFERROR(--RIGHT(B{i},6),LEN(B{i}))*0.0071))*43758.5453,1),"")')
    wh[f"E{i}"] = (f'=IF(C{i}=1,INT(MOD(ABS(SIN(D{i}*7919.123+ROW()*0.618+{SEED}*1.3))*43758.5453,1)*1000000)*10000+ROW(),"")')
    wh[f"F{i}"] = f'=IF(E{i}="","",RANK(E{i},$E$2:$E${LAST},1))'
    wh[f"G{i}"] = f'=IF(A{i}="",0,IF(SUMPRODUCT(--($A$2:$A${LAST}=A{i}))>1,1,0))'
    wh[f"H{i}"] = f"=ROW()-1"
    wh[f"I{i}"] = (f'=IF(OR({R}=0,H{i}>{N}),"",IF(H{i}-1<{REM}*({Q}+1),INT((H{i}-1)/({Q}+1)),'
                   f'{REM}+INT((H{i}-1-{REM}*({Q}+1))/{Q})))')
    wh[f"J{i}"] = (f'=IF(I{i}="","",IF(H{i}-1<{REM}*({Q}+1),MOD(H{i}-1,{Q}+1),MOD(H{i}-1-{REM}*({Q}+1),{Q})))')
    wh[f"K{i}"] = f'=IF(I{i}="","",J{i}*{R}+I{i}+1)'
    wh[f"L{i}"] = f'=IF(K{i}="","",MATCH(K{i},$F$2:$F${LAST},0))'
    wh[f"M{i}"] = f'=IF(I{i}="","",INDEX({S}!$A$2:$A$21,MOD(I{i}+{OFF},{R})+1))'
for col in "ABCDEFGHIJKLM":
    wh.column_dimensions[col].width = 15
wh.sheet_state = "hidden"

# ---------------- תוצאה ----------------
for j, t in enumerate(["מספר זיהוי", "טלפון נייד", "שם נציג"], 1):
    hdr(wr.cell(row=1, column=j), t)
for i in range(2, LAST + 1):
    wr[f"A{i}"] = f'=IF({H}!L{i}="","",INDEX({H}!$A$2:$A${LAST},{H}!L{i}))'
    wr[f"B{i}"] = f'=IF({H}!L{i}="","",INDEX({H}!$B$2:$B${LAST},{H}!L{i}))'
    wr[f"C{i}"] = f'={H}!M{i}'
    for col in "ABC":
        c = wr[f"{col}{i}"]
        c.number_format = "@"
        c.font = font
        c.alignment = center
wr.column_dimensions["A"].width = 18
wr.column_dimensions["B"].width = 18
wr.column_dimensions["C"].width = 16
wr.freeze_panes = "A2"
wr.auto_filter.ref = f"A1:C{LAST}"
band = PatternFill("solid", fgColor="E2EFDA")
wr.conditional_formatting.add(
    f"A2:C{LAST}",
    FormulaRule(formula=[f'AND($C2<>"",MOD(IFERROR(MATCH($C2,{S}!$A$2:$A$21,0),0),2)=0)'], fill=band))
wr.conditional_formatting.add(
    f"A2:C{LAST}",
    FormulaRule(formula=['AND($C2<>"",$C2<>$C1,ROW()>2)'], border=Border(top=Side(style="medium", color="1F4E78"))))

# ---------------- הדבקת לידים ----------------
hdr(wp["A1"], "מספר זיהוי")
hdr(wp["B1"], "טלפון נייד")
for i in range(2, LAST + 1):
    for col in "AB":
        c = wp[f"{col}{i}"]
        c.number_format = "@"
        c.fill = input_fill
        c.font = font
        c.alignment = center
wp.freeze_panes = "A2"
wp.column_dimensions["A"].width = 18
wp.column_dimensions["B"].width = 18
wp.column_dimensions["C"].width = 3
wp.column_dimensions["D"].width = 3
wp.column_dimensions["E"].width = 34
wp.column_dimensions["F"].width = 22
wp.column_dimensions["G"].width = 3
wp.column_dimensions["H"].width = 80


def lab(cell, text, fill=panel_fill, f=bold):
    c = wp[cell]
    c.value = text
    c.font = f
    c.fill = fill
    c.border = box
    c.alignment = right


def val(cell, v, fill=gray_fill):
    c = wp[cell]
    c.value = v
    c.font = bold
    c.fill = fill
    c.border = box
    c.alignment = center


hdr(wp["E1"], "הגרלה")
wp["F1"].fill = hdr_fill
wp["F1"].border = box
lab("E2", "מספר הגרלה (לשינוי הערבוב)")
val("F2", 1, input_fill)
wp["F2"].font = Font(name=F, size=14, bold=True, color="C00000")
dv_seed = DataValidation(type="whole", operator="between", formula1="1", formula2="999999",
                         showErrorMessage=True, errorTitle="ערך לא תקין", error="יש להזין מספר שלם בין 1 ל-999999")
wp.add_data_validation(dv_seed)
dv_seed.add("F2")

hdr(wp["E4"], "בדיקות ובקרה")
wp["F4"].fill = hdr_fill
wp["F4"].border = box
PA = f"$A$2:$A${LAST}"
PB = f"$B$2:$B${LAST}"
checks = [
    (5, "לידים שזוהו", f"={N}"),
    (6, "לידים שחולקו בגיליון התוצאה", f'=COUNTIF({RES}!$C$2:$C${LAST},"?*")'),
    (7, "שורות עם שדה אחד חסר", f'=SUMPRODUCT(--((({PA}<>"")+({PB}<>""))=1))'),
    (8, "מספרי זיהוי שמופיעים יותר מפעם אחת", f"=SUM({H}!$G$2:$G${LAST})"),
    (9, "שורות מעבר לקיבולת (5,000)", f'=COUNTA($A${LAST + 1}:$B$1048576)'),
    (10, "הפרש בין הנציג הגדול לקטן", '=IF(COUNT(F15:F34)=0,0,MAX(F15:F34)-MIN(F15:F34))'),
    (11, "מצב", '=IF(F5=0,"ממתין להדבקת לידים",IF(F9>0,"⚠ יותר מ-5,000 לידים – לא כולם חולקו",'
                'IF(OR(F6<>F5,F10>1),"⚠ שגיאה בחלוקה – בדקו את ההגדרות",'
                'IF(F8>0,"✔ חולק – שימו לב: יש מספרי זיהוי כפולים","✔ החלוקה תקינה"))))'),
]
for r, label, f in checks:
    lab(f"E{r}", label, fill=PatternFill(fill_type=None), f=font)
    val(f"F{r}", f)
wp["F11"].font = Font(name=F, size=11, bold=True, color="006100")
wp.conditional_formatting.add("F11", FormulaRule(formula=['LEFT($F$11,1)="⚠"'],
                                                 font=Font(color="C00000", bold=True),
                                                 fill=PatternFill("solid", fgColor="FFC7CE")))
wp.conditional_formatting.add("F7:F9", FormulaRule(formula=["F7>0"], font=Font(color="C00000", bold=True)))
wp.conditional_formatting.add("F10", FormulaRule(formula=["F10>1"], font=Font(color="C00000", bold=True)))

hdr(wp["E14"], "נציג")
hdr(wp["F14"], "מספר לידים")
for k in range(REP_SLOTS):
    r = 15 + k
    wp[f"E{r}"] = f'=IF({S}!A{k + 2}="","",{S}!A{k + 2})'
    wp[f"F{r}"] = f'=IF(E{r}="","",COUNTIF({RES}!$C$2:$C${LAST},E{r}))'
    for col in "EF":
        wp[f"{col}{r}"].font = font
        wp[f"{col}{r}"].alignment = center
        wp[f"{col}{r}"].border = box
lab("E35", "סה\"כ", fill=panel_fill)
val("F35", "=SUM(F15:F34)")

hdr(wp["H1"], "הוראות שימוש")
instructions = [
    "1.  מחקו את הרשימה הקודמת: סמנו את עמודות A:B (מהשורה 2 ומטה) ולחצו Delete.",
    "2.  העתיקו את הלידים (שתי עמודות בלבד: מספר זיהוי, טלפון נייד – ללא שורת כותרות).",
    "3.  לחצו על תא A2 והדביקו. מומלץ: לחצן ימני ← \"הדבק ערכים\" (או Ctrl+Shift+V) כדי לשמור על העיצוב.",
    "4.  זהו! עברו לגיליון \"תוצאה\" – כל הלידים מעורבבים ומחולקים לנציגים, מקובצים לפי נציג.",
    "5.  ודאו שבתא F11 מופיע \"✔ החלוקה תקינה\" ושההפרש בין הנציגים (F10) הוא 0 או 1.",
    "",
    "קיבוע התוצאה:  הערבוב אינו משתמש ב-RAND ולכן אינו משתנה בפתיחת הקובץ או בחישוב מחדש.",
    "   הוא נקבע לפי רשימת הלידים + \"מספר ההגרלה\" (F2). אותה רשימה ואותו מספר = אותה חלוקה תמיד.",
    "   לערבוב חדש – שנו את מספר ההגרלה (למשל 2, 3...). כדאי להשתמש במספר אחר בכל חודש.",
    "   לשמירת עותק סופי: שמרו את הקובץ בשם חדש (למשל \"חלוקה – ספטמבר\"), או העתיקו את גיליון",
    "   התוצאה והדביקו כ\"ערכים\" בקובץ נפרד לשליחה לנציגים.",
    "",
    "אפסים מובילים:  עמודות ההדבקה מעוצבות כטקסט. אם בכל זאת הודבק מספר שאיבד את האפס,",
    "   הקובץ משלים אותו אוטומטית (ת\"ז ל-9 ספרות, נייד ל-10 ספרות) – ניתן לשנות בגיליון \"הגדרות\".",
    "",
    "קיבולת: עד 5,000 לידים. שמות הנציגים ניתנים לעריכה בגיליון \"הגדרות\" (עד 20 נציגים).",
    "תאים בצהוב = מקום להקלדה/הדבקה. אין צורך למיין, לגרור נוסחאות או להפעיל מאקרו.",
]
for k, t in enumerate(instructions):
    c = wp.cell(row=2 + k, column=8, value=t)
    c.font = Font(name=F, size=11, bold=t[:2] in ("קי", "אפ") or t.startswith("קיבולת"))
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.fill = PatternFill("solid", fgColor="FFFFFF" if t == "" else "F7F9FC")

wb.calculation = CalcProperties(fullCalcOnLoad=True)

# optional test data
if len(sys.argv) > 2:
    import json
    rows = json.load(open(sys.argv[2]))
    for k, (idv, ph) in enumerate(rows):
        wp.cell(row=2 + k, column=1, value=idv)
        wp.cell(row=2 + k, column=2, value=ph)
    if len(sys.argv) > 3:
        wp["F2"] = int(sys.argv[3])

wb.save(OUT)
print("saved", OUT)
